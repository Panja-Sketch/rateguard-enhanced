"""Sheet/column contract enforcement and non-executing openpyxl loading
(locked doc sections 5.1 and 5.3 steps 4-5).

`load_workbook_safely` loads only in non-executing, formula-preserving mode
(`data_only=False`, `keep_vba=False`) -- never `data_only=True`, which would
silently trust Excel's last-cached calculated values instead of the raw
formula/value the compiler must itself validate and, for the mini-DSL sheets,
interpret.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.ingestion.workbook_v1.errors import WorkbookErrorDetail, WorkbookRejectionError

REQUIRED_SHEETS: dict[str, list[str]] = {
    "RG_METADATA": ["key", "value"],
    "RG_INPUTS": [
        "id", "name", "data_type", "required", "minimum", "maximum", "allowed_values",
    ],
    "RG_CONSTANTS": ["id", "name", "value"],
    "RG_TABLES": [
        "table_id", "dimension_id", "min", "max", "include_min", "include_max",
        "match_value", "result_value", "priority",
    ],
    "RG_CALCULATIONS": ["node_id", "operator", "operand_1", "operand_2", "rounding_mode", "scale"],
    "RG_OUTPUTS": ["output_id", "source_ref", "currency"],
    "RG_CONTROL_CASES": ["case_id", "input", "expected_output", "tolerance"],
}


def load_workbook_safely(content: bytes) -> Workbook:
    """Loads a workbook that has already passed `zip_safety.validate_archive_safety`.
    Non-executing, formula-preserving mode: `data_only=False` (raw formula
    strings/values, not Excel's cached calculated results), `keep_vba=False`
    (never load macro content even defensively, though macro-enabled files
    are already rejected upstream)."""
    try:
        return openpyxl.load_workbook(
            io.BytesIO(content), data_only=False, read_only=False, keep_vba=False
        )
    except Exception as exc:  # noqa: BLE001 - any openpyxl failure is a hard rejection
        # User-facing message stays free of raw exception class names/internal
        # archive-member paths (e.g. "BadZipFile: Bad CRC-32 for file
        # 'xl/theme/theme1.xml'"); the technical detail is preserved in
        # `details[].note` for support/debugging, not the headline message.
        raise WorkbookRejectionError(
            code="CORRUPT_ARCHIVE",
            message=(
                "This file isn't a valid Excel workbook -- it may be corrupted, "
                "in an unsupported format, or damaged in transit. Try re-saving "
                "it from Excel and uploading again."
            ),
            details=[WorkbookErrorDetail(note=f"{type(exc).__name__}: {exc}")],
        ) from exc


def _header_index(sheet: Worksheet) -> dict[str, int]:
    """Maps a lowercased header-cell string to its 0-based column index,
    reading only the first row."""
    headers: dict[str, int] = {}
    for idx, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1))):
        if cell.value is not None:
            headers[str(cell.value).strip().lower()] = idx
    return headers


def validate_required_sheets_and_columns(wb: Workbook) -> dict[str, dict[str, int]]:
    """Checks every required sheet exists and every required column header
    is present on it. Returns, per sheet, the header-name -> column-index
    map for use by `extract_rows`."""
    header_maps: dict[str, dict[str, int]] = {}
    for sheet_name, required_columns in REQUIRED_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            raise WorkbookRejectionError(
                code="MISSING_REQUIRED_SHEET",
                message=f"Required sheet '{sheet_name}' is missing from the workbook.",
                details=[WorkbookErrorDetail(sheet=sheet_name)],
            )
        sheet = wb[sheet_name]
        headers = _header_index(sheet)
        header_maps[sheet_name] = headers
        for column in required_columns:
            if column not in headers:
                raise WorkbookRejectionError(
                    code="MISSING_REQUIRED_COLUMN",
                    message=f"Sheet '{sheet_name}' is missing required column '{column}'.",
                    details=[WorkbookErrorDetail(sheet=sheet_name, cell="row1", note=column)],
                )
    return header_maps


def extract_rows(
    wb: Workbook, sheet_name: str, headers: dict[str, int]
) -> list[dict[str, Any]]:
    """Reads every non-blank data row (row 2 onward) of `sheet_name` into a
    dict keyed by the lowercased header names discovered in `headers`."""
    sheet = wb[sheet_name]
    rows: list[dict[str, Any]] = []
    for row_idx, raw_row in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = [cell.value for cell in raw_row]
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in values):
            continue
        row = {name: (values[idx] if idx < len(values) else None) for name, idx in headers.items()}
        row["_row_number"] = row_idx
        rows.append(row)
    return rows


def all_used_cells(wb: Workbook, sheet_names: list[str]):
    """Yields (sheet_name, cell_coordinate, value) for every non-empty cell
    across the given sheets -- used by the raw-formula/unsupported-function
    scan. Restricted to the required sheets (rather than every sheet in the
    workbook) since those are the only sheets this contract reads or
    executes any logic from; an unused extra sheet's content is inert."""
    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    yield sheet_name, cell.coordinate, cell.value
